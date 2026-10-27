"""One-off: dump Proposal docx text for review. Read-only; prints to stdout."""
import glob

import docx

path = glob.glob("docs/*.docx")[0]
print("FILE:", path)
d = docx.Document(path)
print("PARAS:", len(d.paragraphs), "TABLES:", len(d.tables), "INLINE_SHAPES:", len(d.inline_shapes))
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t:
        print("[%d][%s] %s" % (i, para.style.name, t))
for ti, tbl in enumerate(d.tables):
    print("=== TABLE %d: %d rows x %d cols" % (ti, len(tbl.rows), len(tbl.columns)))
    for r in tbl.rows:
        print(" | ".join(c.text.strip().replace("\n", " / ") for c in r.cells))
