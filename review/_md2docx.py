"""Minimal Markdown -> DOCX converter for the review report.

Supports: #/##/### headings, paragraphs, "- " bullets, "1. " numbered items,
pipe tables (with separator row), fenced code blocks, **bold**, `code`,
and a cover block delimited by --- front matter fields title/subtitle/meta.

Usage:
    python _md2docx.py input.md output.docx
"""

from __future__ import annotations

import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "微软雅黑"
EN_FONT = "Segoe UI"
ACCENT = "1F4E79"
GRAY = "595959"


def _set_run_font(run, size=None, bold=None, color=None, font=None):
    font = font or CN_FONT
    run.font.name = EN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _cell_text(cell, text: str, bold=False, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    _add_inline(p, text, size=size, bold=bold)


def _add_inline(paragraph, text: str, size=10.5, bold=False):
    """Render **bold** and `code` spans."""
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            run = paragraph.add_run(tok[2:-2])
            _set_run_font(run, size=size, bold=True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            run = paragraph.add_run(tok[1:-1])
            _set_run_font(run, size=size - 0.5, bold=bold, font="Consolas")
            run.font.color.rgb = RGBColor.from_string("A31F34")
        else:
            run = paragraph.add_run(tok)
            _set_run_font(run, size=size, bold=bold)


def _add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("第 ")
    _set_run_font(run, size=8.5, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    run2 = p.add_run(" 页")
    _set_run_font(run2, size=8.5, color=GRAY)


def _add_cover(doc: Document, title: str, subtitle: str, meta_lines: list[str]):
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    _set_run_font(run, size=26, bold=True, color=ACCENT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(subtitle)
    _set_run_font(run, size=14, bold=False, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    run = p.add_run("─" * 26)
    _set_run_font(run, size=10, color="BFBFBF")

    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        _add_inline(p, line, size=10.5, bold=False)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("（本文件由评审分析生成，供组内讨论与提交参考）")
    _set_run_font(run, size=9, color="808080")
    doc.add_page_break()


def _style_heading(doc: Document):
    styles = doc.styles
    h1 = styles["Heading 1"]
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(ACCENT)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(8)
    h1.font.name = EN_FONT
    h1.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    h2 = styles["Heading 2"]
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string("2E74B5")
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    h2.font.name = EN_FONT
    h2.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    h3 = styles["Heading 3"]
    h3.font.size = Pt(11.5)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string("404040")
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)
    h3.font.name = EN_FONT
    h3.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    normal = styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = EN_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)


def _parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def _set_table_widths(table, rows: list[list[str]], total_cm: float = 16.4):
    def vlen(s: str) -> int:
        return sum(2 if ord(ch) > 127 else 1 for ch in s)

    ncols = max(len(r) for r in rows)
    weights = []
    for ci in range(ncols):
        longest = max([vlen(r[ci]) for r in rows if ci < len(r)] or [4])
        weights.append(min(max(longest, 6), 46))
    total_w = sum(weights) or 1
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for row in table.rows:
        for ci, cell in enumerate(row.cells):
            if ci < len(weights):
                cell.width = Cm(total_cm * weights[ci] / total_w)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _no_row_split(table) -> None:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)


def _meta_block_items(lines: list[str], start: int) -> list[str]:
    """Collect the entries of a block-style `meta:` key (list items or key/value pairs)."""
    items: list[str] = []
    j = start + 1
    while j < len(lines) and lines[j].strip() != "---":
        item = re.match(r"^\s+-\s+(.*)$", lines[j])
        pair = re.match(r"^\s{2,}([^:\s][^:]*?)\s*:\s*(.+)$", lines[j])
        if item:
            items.append(item.group(1).strip())
        elif pair:
            items.append(f"{pair.group(1)}：{pair.group(2).strip()}")
        elif lines[j].strip():
            break
        j += 1
    return items


def convert(md_path: str, out_path: str) -> None:
    raw = open(md_path, encoding="utf-8").read()
    lines = raw.splitlines()

    # --- front matter (title / subtitle / meta) ---
    title = "评审报告"
    subtitle = ""
    meta: list[str] = []
    i = 0
    if lines and lines[0].strip() == "---":
        i = 1
        while i < len(lines) and lines[i].strip() != "---":
            m = re.match(r"^(title|subtitle|meta)\s*:\s*(.*)$", lines[i])
            if m:
                key, val = m.group(1), m.group(2)
                if key == "title":
                    title = val
                elif key == "subtitle":
                    subtitle = val
                elif key == "meta":
                    meta.append(val)
            i += 1
        i += 1

    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    _style_heading(doc)
    _add_page_number_footer(section)

    _add_cover(doc, title, subtitle, meta)

    # --- body ---
    in_code = False
    code_buf: list[str] = []
    table_buf: list[str] = []
    para_buf: list[str] = []

    def flush_para():
        nonlocal para_buf
        if not para_buf:
            return
        text = " ".join(para_buf).strip()
        para_buf = []
        if not text:
            return
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = None
        _add_inline(p, text, size=10.5)

    def flush_table():
        nonlocal table_buf
        rows = _parse_table(table_buf)
        table_buf = []
        if not rows:
            return
        ncols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=ncols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for ri, row in enumerate(rows):
            for ci in range(ncols):
                cell = table.cell(ri, ci)
                text = row[ci] if ci < len(row) else ""
                if ri == 0:
                    _shade(cell, "DEEAF6")
                    _cell_text(cell, text, bold=True)
                else:
                    _cell_text(cell, text)
        _set_table_widths(table, rows)
        _repeat_header(table.rows[0])
        _no_row_split(table)
        # spacing after table
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def flush_code():
        nonlocal code_buf
        if not code_buf:
            return
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run("\n".join(code_buf))
        _set_run_font(run, size=8.5, font="Consolas", color="333333")
        code_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                flush_para()
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if stripped.startswith("|"):
            flush_para()
            table_buf.append(stripped)
            i += 1
            continue
        else:
            if table_buf:
                flush_table()

        if not stripped:
            flush_para()
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            flush_para()
            level = len(m.group(1))
            text = m.group(2).strip().replace("`", "").replace("**", "")
            h = doc.add_heading(level=level)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(text)
            _set_run_font(
                run,
                size={1: 16, 2: 13, 3: 11.5}[level],
                bold=True,
                color={1: ACCENT, 2: "2E74B5", 3: "404040"}[level],
            )
            i += 1
            continue

        mimg = re.match(r"^!\[(.*?)\]\((.+?)\)$", stripped)
        if mimg:
            flush_para()
            flush_table()
            import os

            img_path = mimg.group(2)
            if not os.path.isabs(img_path):
                img_path = os.path.join(os.path.dirname(md_path), img_path)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(img_path, width=Cm(16.2))
            if mimg.group(1):
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(10)
                r = cap.add_run("图 " + mimg.group(1))
                _set_run_font(r, size=8.5, color=GRAY)
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            flush_para()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            _add_inline(p, re.sub(r"^[-*]\s+", "", stripped), size=10.5)
            i += 1
            continue

        mnum = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if mnum:
            flush_para()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.85)
            p.paragraph_format.first_line_indent = Cm(-0.85)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            _add_inline(p, "%s. %s" % (mnum.group(1), mnum.group(2)), size=10.5)
            i += 1
            continue

        if stripped.startswith(">"):
            flush_para()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped.lstrip("> ").strip())
            _set_run_font(run, size=10, color=GRAY)
            i += 1
            continue

        if stripped == "---":
            flush_para()
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("─" * 40)
            _set_run_font(run, size=8, color="D9D9D9")
            i += 1
            continue

        para_buf.append(stripped)
        i += 1

    flush_para()
    flush_table()
    if in_code:
        flush_code()

    cp = doc.core_properties
    cp.title = title
    cp.subject = subtitle
    cp.author = "MAIE 6000C Project Group 15"
    cp.comments = "Proposal review, architecture recommendations, four-person division of work"

    doc.save(out_path)
    print("saved:", out_path)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
