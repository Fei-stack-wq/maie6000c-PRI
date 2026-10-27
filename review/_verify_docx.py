"""Quick validation of the generated DOCX (read-only)."""

import glob

import docx

path = glob.glob("review/*.docx")[0]
d = docx.Document(path)
print("FILE:", path)
print("PARAS:", len(d.paragraphs), "TABLES:", len(d.tables), "INLINE_SHAPES:", len(d.inline_shapes))
print("TITLE_PROP:", d.core_properties.title)
heads = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
print("H1:", " | ".join(heads))
text = "\n".join(p.text for p in d.paragraphs)
for key in ["P0-10", "R1 数据与接入", "结论摘要", "8 个决策", "修订后状态机（二）"]:
    print(key, "->", key in text)
